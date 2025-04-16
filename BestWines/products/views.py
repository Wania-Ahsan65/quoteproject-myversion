from django.shortcuts import render, redirect, get_list_or_404
from django.contrib import messages
from openpyxl import load_workbook
from .forms import ProductUploadForm
from .models import Product
from django.core.paginator import Paginator
from django.db.models import Q
from collections import defaultdict
from django.forms import modelformset_factory
from .forms import GeneralProductForm, VariantProductForm

def product_upload_view(request):
    if request.method == 'POST':
        form = ProductUploadForm(request.POST, request.FILES)
        if form.is_valid():
            file = request.FILES['file']
            wb = load_workbook(file)
            sheet = wb.active

            # Process each row (skipping header row)
            new_count = 0
            skipped_count = 0
            for row in sheet.iter_rows(min_row=2, values_only=True):
                product_code = row[0]
                product_description = row[1]
                category = row[2]
                size = row[3]
                abv = row[4]
                country = row[5]
                brand = row[6]
                cost_price = row[7]

                if product_code:  # Ensure product_code is present
                    # Check for a product with EXACTLY these values
                    exists = Product.objects.filter(
                        product_code=product_code,
                        product_description=product_description,
                        category=category,
                        size=size,
                        abv=abv,
                        country=country,
                        brand=brand,
                        cost_price=cost_price,
                    ).exists()

                    if exists:
                        skipped_count += 1
                        continue
                    else:
                        Product.objects.create(
                            product_code=product_code,
                            product_description=product_description,
                            category=category,
                            size=size,
                            abv=abv,
                            country=country,
                            brand=brand,
                            cost_price=cost_price,
                        )
                        new_count += 1

            messages.success(request,
                f"Products uploaded. {new_count} new records added; {skipped_count} duplicates skipped.")
            return redirect('products:list')
    else:
        form = ProductUploadForm()
    
    return render(request, 'products/upload.html', {'form': form})

def product_list_view(request):
    query = request.GET.get('q', '')
    if query:
        products_qs = Product.objects.filter(
            Q(product_code__icontains=query) | Q(product_description__icontains=query)
        ).order_by('product_code')
    else:
        products_qs = Product.objects.all().order_by('product_code')
    
    # Group products by product_code
    grouped = defaultdict(list)
    for product in products_qs:
        grouped[product.product_code].append(product)
    
    grouped_products = []
    for product_code, variants in grouped.items():
        main_info = variants[0]  # Use first record as "main" product info
        grouped_products.append({
            'product_code': product_code,
            'product_description': main_info.product_description,
            'category': main_info.category,
            'size': main_info.size,
            'abv': main_info.abv,
            'country': main_info.country,
            'variants': [{'brand': variant.brand, 'cost_price': variant.cost_price} for variant in variants],
        })
    
    # Paginate the grouped products (10 groups per page)
    paginator = Paginator(grouped_products, 10)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    return render(request, 'products/product_list.html', {
        'page_obj': page_obj,
        'query': query,
    })


from django.shortcuts import render, redirect
from django.forms import modelformset_factory
from django.contrib import messages
from .models import Product
from .forms import GeneralProductForm, VariantProductForm

def product_edit_view(request, product_code):
    queryset = Product.objects.filter(product_code=product_code)
    if not queryset.exists():
        messages.error(request, "Product group not found.")
        return redirect('products:list')

    common_instance = queryset.first()
    VariantFormSet = modelformset_factory(Product, form=VariantProductForm, extra=0)
    
    if request.method == 'POST':
        general_form = GeneralProductForm(request.POST, instance=common_instance)
        variant_formset = VariantFormSet(request.POST, queryset=queryset)
        if general_form.is_valid() and variant_formset.is_valid():
            general_instance = general_form.save(commit=False)
            queryset.update(
                product_description=general_instance.product_description,
                category=general_instance.category,
                size=general_instance.size,
                abv=general_instance.abv,
                country=general_instance.country,
            )
            variant_formset.save()
            messages.success(request, "Changes applied successfully. Price updated!")
            return redirect('products:list')
        else:
            print("General Form Errors:", general_form.errors)
            print("Variant Formset Errors:", variant_formset.errors)
            messages.error(request, "There were errors in your submission. Please fix them below.")
    else:
        general_form = GeneralProductForm(instance=common_instance)
        variant_formset = VariantFormSet(queryset=queryset)
    
    context = {
        'general_form': general_form,
        'variant_formset': variant_formset,
        'product_code': product_code,
    }
    return render(request, 'products/product_edit_new.html', context)
