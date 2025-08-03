from django.shortcuts import render, get_object_or_404, redirect
from .models import Quote, QuoteItem
from django.contrib.auth.decorators import login_required
from django.forms import modelform_factory
from django.forms import inlineformset_factory
from products.models import Product
from .models import Quote, QuoteItem
from customers.models import Customer
from django.http import HttpResponse
from django.template.loader import render_to_string
from decimal import Decimal

@login_required
def quote_search_view(request):
    query = request.GET.get("q", "")
    if request.user.is_superuser:
        quotes = Quote.objects.all()
    else:
        quotes = Quote.objects.filter(user=request.user)

    if query:
        quotes = quotes.filter(customer_name__icontains=query)

    html = render_to_string("quotes/quote_list_partial.html", {"quotes": quotes})
    return HttpResponse(html)

from django.shortcuts import render

def quote_list_view(request):
    # Replace with your actual logic
    return render(request, 'quotes/quote_list.html')

from django.shortcuts import render, redirect
from .models import Quote, QuoteItem
from products.models import Product
from collections import defaultdict

def quote_create_view(request):

    customers = Customer.objects.all()  

    if request.method == 'POST':
        discount=Decimal(request.POST.get("discount") or 0)

        quote = Quote.objects.create(
            customer_name=request.POST.get("customer_name"),
            notes=request.POST.get("notes"),
            user=request.user,
            discount=discount
        )
       
        items = request.POST.getlist("items[]")  # or parse from `items[product_id]` format

        # OR loop manually if you're using dynamic field names like items[123][quantity]
        for key in request.POST:
            if key.startswith('items[') and key.endswith('][product_id]'):
                prefix = key.split('][')[0]  # e.g., items[123
                product_id = request.POST.get(key)
                quantity = request.POST.get(f"{prefix}][quantity]")
                cost = Decimal(request.POST.get(f"{prefix}][cost_price]"))
                selling = Decimal(request.POST.get(f"{prefix}][selling_price]"))

                margin = ((selling - cost) / selling ) * 100 if selling else 0
                discount_price = selling - (selling * discount/100)


                if product_id and quantity:
                    QuoteItem.objects.create(
                        quote=quote,
                        product_id=product_id,
                        quantity=quantity,
                        cost_price=cost,
                        selling_price=selling,
                        margin=margin,
                        discount=discount,
                        discount_price=discount_price     
                    )

        return redirect('quotes:detail', pk=quote.pk)

    # GET request
    grouped = defaultdict(list)
    for product in Product.objects.all():
        grouped[product.product_code].append(product)


    categories = Product.objects.values_list('category', flat=True).distinct()
    return render(request, 'quotes/quote_create.html', {
        'products_by_code': dict(grouped),
        'categories': categories,
        'customers': customers,

    })

from django.contrib.auth.decorators import login_required
from django.shortcuts import get_object_or_404, render
from .models import Quote, QuoteItem

@login_required
def quote_detail_view(request, pk):
    if request.user.is_superuser:
        quote = get_object_or_404(Quote, pk=pk)
    else:
        quote = get_object_or_404(Quote, pk=pk, user=request.user)

    if request.method == 'POST':
        for key in request.POST:
            if key.startswith('item_') and key.endswith('_id'):
                prefix = key.replace('_id', '')
                item_id = request.POST.get(f"{prefix}_id")
                item = QuoteItem.objects.get(pk=item_id, quote=quote)

                item.quantity = 1
                cost_price = request.POST.get(f"{prefix}_cost_price")
                selling_price = request.POST.get(f"{prefix}_selling_price")
                margin = request.POST.get(f"{prefix}_margin")
                discount = request.POST.get(f"{prefix}_discount")  # Expecting a percentage input

                # Convert values to Decimal for precise arithmetic
                from decimal import Decimal

                item.cost_price = Decimal(cost_price or 0)
                item.selling_price = Decimal(selling_price or 0)
                item.margin = Decimal(margin or 0)
                item.discount = Decimal(discount or 0)

                # Calculate discounted price
                item.discount_price = item.selling_price - (item.selling_price * item.discount / Decimal('100'))

                item.save()

        return redirect('quotes:detail', pk=quote.pk)

    return render(request, 'quotes/quote_detail.html', {'quote': quote})


from io import BytesIO
from django.http import HttpResponse
import openpyxl
from openpyxl.utils import get_column_letter
from .models import Quote

def export_quote_full(request, pk):
    quote = Quote.objects.get(pk=pk)
    items = quote.items.select_related('product')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Quote #{quote.id} - Full"

    # Header
    headers = [
        "Code", "Product Name", "Category", "Size", "ABV", "Country", "Brand",
        "Cost Price", "Selling Price", "Margin (%)"
    ]
    ws.append(headers)

    for item in items:
        product = item.product

        # Recalculate margin for safety
        try:
            cost = float(item.cost_price)
            price = float(item.selling_price)
            margin = ((price - cost) / price * 100) if price > 0 else 0
        except:
            margin = 0

        row = [
            product.product_code,
            product.product_description,
            product.category,
            product.size,
            product.abv,
            product.country,
            product.brand,
            cost,
            price,
            round(margin, 2)
        ]
        ws.append(row)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        content=output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=Quote-{quote.id}-Full.xlsx'
    return response




from io import BytesIO
from django.http import HttpResponse
import openpyxl
from openpyxl.utils import get_column_letter
from .models import Quote

def export_quote_sales(request, pk):
    quote = Quote.objects.get(pk=pk)
    items = quote.items.select_related('product')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Quote #{quote.id} - Client"

    # ✅ New Header Format
    headers = ["Product Code", "Description", "Size", "Price"]
    ws.append(headers)

    for item in items:
        product = item.product
        row = [
            product.product_code,
            product.product_description,
            product.size,
            float(item.selling_price)
        ]
        ws.append(row)

    for col in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 20

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    response = HttpResponse(
        content=output,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename=Quote-{quote.id}-Client.xlsx'
    return response


from django.template.loader import get_template
from xhtml2pdf import pisa
from io import BytesIO
from django.http import HttpResponse
from .models import Quote
from django.template.loader import get_template
from django.template import Context
from django.contrib.staticfiles import finders
from django.utils.dateformat import format 
from django.utils.text import Truncator

def export_quote_client_pdf(request, pk):
    quote = Quote.objects.get(pk=pk)
    items = quote.items.select_related('product')

    formatted_date = format(quote.created_at, "d M Y, h:i A")

    for item in items:
        desc = item.product.product_description
        clean_desc = desc.split(' - ')[0]  # Removes part after ' - '
        item.product.cleaned_description = clean_desc

    template = get_template('quotes/client_quote_pdf.html')
    html = template.render({'quote': quote, 'items': items, 'created_at': formatted_date})

    response = HttpResponse(content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="Quote-{quote.id}-Client.pdf"'

    pisa_status = pisa.CreatePDF(
        src=html,
        dest=response
    )
    if pisa_status.err:
        return HttpResponse("We had some errors with PDF generation", status=500)
    return response

